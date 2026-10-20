"""Report building and Excel / CSV / PDF export."""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from database.db import Database
from services.data_service import DataFilters, DataService
from services.export_service import (
    CALCULATED_COLUMNS,
    EXPORT_FORMATS,
    ExportRequest,
    column_labels,
    default_export_path,
    display_frame,
    export,
    provenance_frame,
    provenance_rows,
)
from services.formula_service import FormulaService
from services.report_service import REPORT_KINDS, REPORT_TITLES, ReportService
from services.settings_service import SettingsService


@pytest.fixture()
def reports(filled_db: Database, filled_data: DataService) -> ReportService:
    return ReportService(filled_db, filled_data)


@pytest.fixture()
def sample_frame() -> pd.DataFrame:
    return pd.DataFrame([
        {"symbol": "ACME", "trade_date": _dt.date(2025, 1, 1), "close": 105.0,
         "delivery_pct": 40.0, "avg_price": 102.5},
        {"symbol": "BETA", "trade_date": _dt.date(2025, 1, 2), "close": 210.0,
         "delivery_pct": 55.0, "avg_price": 208.0},
    ])


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------


def test_display_frame_turns_timestamps_into_dates() -> None:
    frame = pd.DataFrame({"trade_date": pd.to_datetime(["2025-01-01", "2025-01-02"])})
    out = display_frame(frame)
    assert out["trade_date"].tolist() == [_dt.date(2025, 1, 1), _dt.date(2025, 1, 2)]


def test_column_labels_mark_calculated_columns(sample_frame: pd.DataFrame) -> None:
    labels = column_labels(sample_frame)
    by_column = dict(zip(sample_frame.columns, labels))
    assert by_column["symbol"] == "Symbol"
    assert by_column["delivery_pct"] == "Delivery %"
    assert by_column["avg_price"] == "Avg Price *"
    assert "avg_price" in CALCULATED_COLUMNS
    assert "close" not in CALCULATED_COLUMNS


def test_provenance_block_records_everything(sample_frame: pd.DataFrame) -> None:
    request = ExportRequest(title="Test", frame=sample_frame, path="out.csv",
                            filters={"formula_mode": "original", "symbols": "ACME"},
                            notes="a note")
    rows = dict(provenance_rows(request))
    assert rows["Report"] == "Test"
    assert rows["Formula mode"] == "original"
    assert rows["Filter: symbols"] == "ACME"
    assert "avg_price" in rows["Calculated columns"]
    assert rows["Notes"] == "a note"

    frame = provenance_frame(request)
    assert list(frame.columns) == ["Field", "Value"]


def test_default_export_path_is_timestamped() -> None:
    path = default_export_path("Screener Report", "csv")
    assert path.suffix == ".csv"
    assert path.stem.startswith("screener-report-")
    assert len(path.stem.rsplit("-", 2)) == 3


# --------------------------------------------------------------------------
# format dispatch
# --------------------------------------------------------------------------


def test_export_rejects_unknown_formats(sample_frame: pd.DataFrame, tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        export(ExportRequest(title="x", frame=sample_frame,
                             path=tmp_path / "x.docx", format="docx"))


def test_csv_export_layout(sample_frame: pd.DataFrame, tmp_path: Path) -> None:
    result = export(ExportRequest(title="Screener", frame=sample_frame,
                                  path=tmp_path / "out.csv", format="CSV",
                                  filters={"formula_mode": "original"}))
    assert result.format == "csv"
    assert result.rows == 2
    assert result.path.exists()

    text = result.path.read_text(encoding="utf-8-sig")
    lines = [line for line in text.splitlines()]
    assert "Generated at" in text
    assert "Columns marked with *" in text
    header_index = next(i for i, line in enumerate(lines) if line.startswith("Symbol"))
    assert "Delivery %" in lines[header_index]
    assert "Avg Price *" in lines[header_index]
    assert lines[header_index + 1].startswith("ACME")


def test_csv_export_blanks_missing_values(tmp_path: Path) -> None:
    frame = pd.DataFrame([{"symbol": "ACME", "close": None, "avg_price": 1.0}])
    result = export(ExportRequest(title="Blanks", frame=frame,
                                  path=tmp_path / "blanks.csv", format="csv"))
    rows = result.path.read_text(encoding="utf-8-sig").splitlines()
    data_row = next(row for row in rows if row.startswith("ACME"))
    assert data_row == "ACME,,1.0"


def test_excel_export_has_provenance_and_data_sheets(sample_frame: pd.DataFrame,
                                                     tmp_path: Path) -> None:
    result = export(ExportRequest(title="Dashboard Snapshot", frame=sample_frame,
                                  path=tmp_path / "out.xlsx", format="xlsx",
                                  filters={"formula_mode": "original"}))
    assert result.path.exists()

    workbook = openpyxl.load_workbook(result.path)
    assert "Report Info" in workbook.sheetnames
    assert "Dashboard Snapshot" in workbook.sheetnames

    data = workbook["Dashboard Snapshot"]
    headers = [cell.value for cell in data[1]]
    assert headers[0] == "Symbol"
    assert "Avg Price *" in headers
    assert data["A2"].value == "ACME"
    assert data.freeze_panes == "A2"

    info = workbook["Report Info"]
    assert info["A1"].value == "Dashboard Snapshot"


def test_excel_export_includes_extra_sheets(sample_frame: pd.DataFrame,
                                            tmp_path: Path) -> None:
    extra = pd.DataFrame({"classification": ["DV1"], "stocks": [3]})
    result = export(ExportRequest(title="Report", frame=sample_frame,
                                  path=tmp_path / "out.xlsx", format="xlsx",
                                  extra_sheets={"Classification": extra}))
    workbook = openpyxl.load_workbook(result.path)
    assert "Classification" in workbook.sheetnames
    assert workbook["Classification"]["A1"].value == "Classification"


def test_excel_export_sanitises_sheet_names(sample_frame: pd.DataFrame,
                                            tmp_path: Path) -> None:
    result = export(ExportRequest(title="Bad/Name[with]chars", frame=sample_frame,
                                  path=tmp_path / "out.xlsx", format="xlsx"))
    workbook = openpyxl.load_workbook(result.path)
    assert "BadNamewithchars" in workbook.sheetnames


def test_pdf_export_writes_a_real_pdf(sample_frame: pd.DataFrame,
                                      tmp_path: Path) -> None:
    result = export(ExportRequest(title="Screener", frame=sample_frame,
                                  path=tmp_path / "out.pdf", format="pdf"))
    assert result.path.read_bytes()[:5] == b"%PDF-"
    assert result.path.stat().st_size > 1000


def test_export_formats_are_documented() -> None:
    assert EXPORT_FORMATS == ("xlsx", "csv", "pdf")


def test_export_creates_missing_directories(sample_frame: pd.DataFrame,
                                            tmp_path: Path) -> None:
    target = tmp_path / "deep" / "nested" / "report.csv"
    export(ExportRequest(title="x", frame=sample_frame, path=target, format="csv"))
    assert target.exists()


# --------------------------------------------------------------------------
# report builders
# --------------------------------------------------------------------------


def test_every_report_kind_builds(reports: ReportService) -> None:
    for kind in REPORT_KINDS:
        report = reports.build(kind)
        assert report.kind == kind
        assert report.title == REPORT_TITLES[kind]
        assert report.filters["formula_mode"] == "original"


def test_unknown_report_kind_raises(reports: ReportService) -> None:
    with pytest.raises(ValueError):
        reports.build("everything")


def test_screener_report_has_one_row_per_symbol(reports: ReportService) -> None:
    report = reports.build("screener")
    assert len(report.frame) == 191
    assert "trade_date" not in report.frame.columns
    assert "symbol" in report.frame.columns


def test_screener_report_accepts_a_universe(reports: ReportService) -> None:
    universe = reports.data.symbols()[:3]
    report = reports.build("screener", universe=universe)
    assert len(report.frame) == 3
    assert set(report.frame["symbol"]) == set(universe)


def test_history_report_is_sorted_by_symbol(reports: ReportService) -> None:
    report = reports.build("history")
    frame = report.frame
    assert len(frame) == 191
    assert frame["symbol"].tolist() == sorted(frame["symbol"].tolist())
    assert list(frame.columns)[0] == "symbol"


def test_history_report_can_append_formula_columns(reports: ReportService) -> None:
    plain = reports.build("history")
    enriched = reports.build("history", include_calculated=True)
    assert len(enriched.frame.columns) > len(plain.frame.columns)
    assert any("Formula Lab" in note for note in [enriched.notes])


def test_lookback_report_has_every_window(reports: ReportService) -> None:
    report = reports.build("lookback", DataFilters(symbols=[reports.data.symbols()[0]]))
    assert len(report.frame) == 5
    assert set(report.frame["window"]) == {"1D", "3D", "5D", "7D", "10D"}


def test_classification_report_counts_every_symbol(reports: ReportService) -> None:
    report = reports.build("classification")
    assert set(report.frame["classification"]) == {"D. Ind", "DV1", "DV2"}
    per_group = report.frame.groupby("classification")["stocks"].sum()
    assert set(per_group) == {191}


def test_formulas_report_lists_the_outputs(reports: ReportService) -> None:
    report = reports.build("formulas")
    assert "close_ma_5" in report.frame.columns
    assert "Formula Lab outputs" in report.notes


def test_dashboard_report_snapshots_the_components(reports: ReportService) -> None:
    report = reports.build("dashboard")
    assert not report.frame.empty
    assert set(report.frame.columns) == {"component", "kind", "value", "status"}
    assert "Unique Stocks" in set(report.frame["component"])


def test_reports_respect_a_symbol_filter(reports: ReportService) -> None:
    symbol = reports.data.symbols()[0]
    report = reports.build("screener", DataFilters(symbols=[symbol]))
    assert report.frame["symbol"].tolist() == [symbol]


# --------------------------------------------------------------------------
# report export / preview
# --------------------------------------------------------------------------


def test_export_uses_the_configured_default_format(db: Database,
                                                   tmp_path: Path) -> None:
    settings = SettingsService(db)
    settings.set("export.directory", str(tmp_path))
    settings.set("export.default_format", "csv")
    service = ReportService(db, DataService(db), FormulaService(db), settings)

    result = service.export(service.build("classification"))
    assert result.format == "csv"
    assert result.path.parent == tmp_path
    assert result.path.exists()
    assert result.path.suffix == ".csv"


def test_export_writes_an_audit_entry(reports: ReportService, tmp_path: Path,
                                      filled_db: Database) -> None:
    report = reports.build("classification")
    result = reports.export(report, tmp_path / "logged.xlsx", fmt="xlsx")
    entries = DataService(filled_db).audit_log(category="export")
    assert any(result.path.name in entry["message"] for entry in entries)


def test_export_rejects_an_unknown_format(reports: ReportService,
                                         tmp_path: Path) -> None:
    report = reports.build("classification")
    with pytest.raises(ValueError):
        reports.export(report, tmp_path / "x.docx", fmt="docx")


def test_preview_applies_labels_and_marks_calculated(reports: ReportService) -> None:
    preview = reports.preview(reports.build("screener"), rows=5)
    assert len(preview) == 5
    assert "Symbol" in preview.columns
    # calculated metrics are always flagged so they cannot be mistaken for source data
    assert any(column.endswith(" *") for column in preview.columns)


def test_provision_frame_is_readable(reports: ReportService) -> None:
    frame = reports.provision_frame(reports.build("classification"))
    assert list(frame.columns) == ["Field", "Value"]
    assert "Generated at" in set(frame["Field"])


def test_export_of_an_empty_report_still_writes_a_file(tmp_path: Path) -> None:
    empty = pd.DataFrame(columns=["symbol", "close"])
    result = export(ExportRequest(title="Empty", frame=empty,
                                  path=tmp_path / "empty.csv", format="csv"))
    assert result.rows == 0
    assert result.path.exists()
    assert "Symbol" in result.path.read_text(encoding="utf-8-sig")

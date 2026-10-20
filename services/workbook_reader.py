"""Workbook inspection utilities.

This module reads the supplied ``.xlsm`` workbook in a strictly read-only way:

* it never executes macros (VBA is only *decompressed as text* for reporting),
* it never writes to the source file,
* it exposes structured metadata (sheets, tables, defined names, data
  validation, conditional formatting, comments, cached values) plus the raw
  ``Data`` records.

The heavy lifting lives here so that ``scripts/inspect_workbook.py``,
``scripts/audit_formulas.py``, the importer and the test-suite all share one
implementation.
"""

from __future__ import annotations

import datetime as _dt
import io
import math
import os
import re
import zipfile
from dataclasses import dataclass, field
from itertools import chain
from pathlib import Path
from typing import Any, Dict, Iterator, List, Optional, Sequence, Tuple

import openpyxl

# --------------------------------------------------------------------------
# Canonical schema of the ``Data`` worksheet
# --------------------------------------------------------------------------

DATA_COLUMNS: Tuple[str, ...] = (
    "Symbol",
    "Date",
    "OPEN",
    "HIGH",
    "LOW",
    "CLOSE",
    "LAST",
    "PREVCLOSE",
    "TOTTRDQTY",
    "TOTTRDVAL",
    "TOTALTRADES",
    "Delivery %",
)

DATE_COLUMN = "Date"
SYMBOL_COLUMN = "Symbol"

_HEADER_ALIASES: Dict[str, Tuple[str, ...]] = {
    "Symbol": ("symbol", "scrip", "ticker", "code", "security"),
    "Date": ("date", "trade date", "tradedate", "trading date", "time"),
    "OPEN": ("open", "o", "open price"),
    "HIGH": ("high", "h", "high price"),
    "LOW": ("low", "l", "low price"),
    "CLOSE": ("close", "c", "close price", "closep"),
    "LAST": ("last", "ltp", "last price", "last traded price", "lastprice"),
    "PREVCLOSE": ("prevclose", "prev close", "previous close", "prev_close", "pc"),
    "TOTTRDQTY": ("tottrdqty", "total traded quantity", "volume", "qty", "traded qty"),
    "TOTTRDVAL": ("tottrdval", "total traded value", "value", "turnover", "traded value"),
    "TOTALTRADES": ("totaltrades", "total trades", "trades", "no of trades", "num trades"),
    "Delivery %": ("delivery %", "delivery%", "delivery", "delivery percent", "deliv perc", "delivery_pct"),
}


def normalise_header(value: Any) -> str:
    """Normalise a raw Excel header into a comparable token."""
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.replace("\n", " ").replace("\u00a0", " ")
    text = re.sub(r"[_\-]+", " ", text)
    text = re.sub(r"\s*%\s*", " %", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def match_canonical_column(header: Any, aliases: Optional[Dict[str, Tuple[str, ...]]] = None) -> Optional[str]:
    """Return the canonical capital column name for a raw header, if any."""
    aliases = aliases or _HEADER_ALIASES
    token = normalise_header(header)
    if not token:
        return None
    for canonical, names in aliases.items():
        if token == normalise_header(canonical):
            return canonical
    for canonical, names in aliases.items():
        if token in {normalise_header(n) for n in names}:
            return canonical
    # Looser containment fallback (e.g. "Total Traded Quantity (Qty)").
    # Single-letter and two-letter aliases ("o", "h", "l", "c", "pc") are only
    # trusted as *whole* tokens: without this guard any header merely containing
    # an "h" would be claimed for HIGH.
    for canonical, names in aliases.items():
        for n in names:
            nn = normalise_header(n)
            if len(nn) >= 3 and len(token) >= 3 and (nn in token or token in nn):
                return canonical
    return None


# --------------------------------------------------------------------------
# Value coercion helpers
# --------------------------------------------------------------------------


def coerce_date(value: Any) -> Optional[_dt.date]:
    """Best-effort conversion of an Excel cell into a ``datetime.date``."""
    if value is None or value == "":
        return None
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Excel serial date (1900 date system)
        try:
            return (_dt.datetime(1899, 12, 30) + _dt.timedelta(days=float(value))).date()
        except (OverflowError, ValueError):
            return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%m/%d/%Y",
                "%d-%b-%Y", "%d %b %Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return _dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return _dt.date.fromisoformat(text[:10])
    except ValueError:
        return None


def coerce_number(value: Any) -> Optional[float]:
    """Convert a cell to float, distinguishing blank/missing from numeric zero."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text in ("", "-", "--", "N/A", "NA", "nan", "None"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def coerce_symbol(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip().upper()
    return text or None


# --------------------------------------------------------------------------
# Streaming worksheet access
# --------------------------------------------------------------------------

try:  # optional: the Rust reader is ~300x faster than openpyxl on big sheets
    from python_calamine import CalamineWorkbook as _CalamineWorkbook
    from python_calamine import WorksheetNotFound as _CalamineWorksheetNotFound
except ImportError:  # pragma: no cover - exercised when the extra is absent
    _CalamineWorkbook = None
    _CalamineWorksheetNotFound = ()

#: Sheet names the importer prefers when no sheet is given.
PREFERRED_SHEETS = ("Data", "data", "Sheet1")


def workbook_engine() -> str:
    """Which reader to use: ``calamine`` when installed, otherwise ``openpyxl``.

    Set ``CAPITAL_WORKBOOK_ENGINE=openpyxl`` to force the pure-Python reader (for
    troubleshooting a workbook that the Rust reader mis-reads) or
    ``CAPITAL_WORKBOOK_ENGINE=calamine`` to require it.
    """
    forced = os.environ.get("CAPITAL_WORKBOOK_ENGINE", "").strip().lower()
    if forced == "openpyxl":
        return "openpyxl"
    if forced == "calamine" and _CalamineWorkbook is None:
        raise RuntimeError("python-calamine is not installed")
    return "calamine" if _CalamineWorkbook is not None else "openpyxl"


class _OpenpyxlSheet:
    """Read-only, values-only worksheet adapter (pure Python fallback)."""

    def __init__(self, path: str | Path, sheet: str):
        self._workbook = openpyxl.load_workbook(
            path, read_only=True, data_only=True, keep_vba=False)
        if sheet not in self._workbook.sheetnames:
            self._workbook.close()
            raise ValueError(f"worksheet {sheet!r} not found in {path}")
        self._worksheet = self._workbook[sheet]
        self.max_row = getattr(self._worksheet, "max_row", None)
        self.max_column = getattr(self._worksheet, "max_column", None)

    def rows(self) -> Iterator[Sequence[Any]]:
        return self._worksheet.iter_rows(values_only=True)

    def close(self) -> None:
        self._workbook.close()


class _CalamineSheet:
    """Worksheet adapter backed by the Rust ``python-calamine`` reader.

    Empty cells come back as ``""`` where openpyxl uses ``None``; they are
    normalised here so every downstream rule sees exactly the same values
    regardless of which reader is active.
    """

    def __init__(self, path: str | Path, sheet: str):
        try:
            self._workbook = _CalamineWorkbook.from_path(str(path))
        except OSError as exc:
            if not Path(path).exists():
                raise FileNotFoundError(str(exc)) from exc
            raise
        try:
            self._sheet = self._workbook.get_sheet_by_name(sheet)
        except _CalamineWorksheetNotFound as exc:
            self._workbook.close()
            raise ValueError(f"worksheet {sheet!r} not found in {path}") from exc
        self.max_row = int(self._sheet.height) or None
        self.max_column = int(self._sheet.width) or None

    def rows(self) -> Iterator[Sequence[Any]]:
        for row in self._sheet.iter_rows():
            yield tuple(None if value == "" else value for value in row)

    def close(self) -> None:
        self._workbook.close()


def open_sheet(path: str | Path, sheet: str):
    """Open one worksheet for streaming with the active engine."""
    if workbook_engine() == "calamine":
        return _CalamineSheet(path, sheet)
    return _OpenpyxlSheet(path, sheet)


def sheet_names(path: str | Path) -> List[str]:
    """Sheet names in workbook order, without loading any cell values."""
    if workbook_engine() == "calamine":
        workbook = _CalamineWorkbook.from_path(str(path))
        try:
            return list(workbook.sheet_names)
        finally:
            workbook.close()
    workbook = openpyxl.load_workbook(path, read_only=True, keep_vba=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def pick_sheet(path: str | Path) -> Optional[str]:
    """Pick the sheet that looks like market data when none was specified."""
    names = sheet_names(path)
    for candidate in PREFERRED_SHEETS:
        if candidate in names:
            return candidate
    for name in names:
        source = None
        try:
            source = open_sheet(path, name)
            first_row = next(iter(source.rows()), None)
        except Exception:  # unreadable sheet: try the next one
            continue
        finally:
            if source is not None:
                source.close()
        tokens = {normalise_header(value) for value in (first_row or ())}
        if "symbol" in tokens or "date" in tokens:
            return name
    return names[0] if names else None


def is_blank_row(values: Any) -> bool:
    """True when every cell of a row is empty or whitespace."""
    return all(value is None or str(value).strip() == "" for value in values)


def _header_cells(values: Any) -> List[str]:
    """Render a header row exactly as :meth:`WorkbookInspector.read_data_records` does."""
    return [str(value).strip() if value is not None else f"Column{index}"
            for index, value in enumerate(values, start=1)]


def _fit_row(values: Any, width: int) -> Tuple[Any, ...]:
    """Pad/trim a streamed row so ``values[i]`` always matches ``headers[i]``."""
    values = tuple(values)
    if len(values) == width:
        return values
    if len(values) > width:
        return values[:width]
    return values + (None,) * (width - len(values))


class DataStream:
    """Lazily read one worksheet of a workbook (read-only, values-only).

    The sheet is streamed, so a worksheet with a million rows costs a constant
    amount of memory — unlike :func:`openpyxl.load_workbook` with
    ``read_only=False``, which builds a Python object for every cell first.
    (The previous implementation also called ``ws.max_column`` inside its row
    loop; that property walks every cell of the sheet, which made a 20k-row read
    take minutes.)

    Reading goes through ``python-calamine`` when it is installed and falls back
    to ``openpyxl`` otherwise; both produce identical values.

    Use it as a context manager::

        with DataStream(path, sheet="Data") as stream:
            headers = stream.headers
            for row_number, values in stream.rows():
                ...

    ``headers`` is the header row rendered as text (blank cells become
    ``ColumnN``) and ``rows()`` yields ``(row_number, values)`` for every row
    after it, padded to the header width.  ``records()`` wraps that in the
    dict-per-row shape the validators expect.
    """

    def __init__(self, path: str | Path, sheet: Optional[str] = None,
                 header_row: Optional[int] = None, scan_rows: int = 25):
        self.path = Path(path)
        self.sheet = sheet
        self.header_row: Optional[int] = int(header_row) if header_row else None
        self.scan_rows = max(1, int(scan_rows))
        self.headers: List[str] = []
        self.width = 0
        self.max_row: Optional[int] = None
        self.max_column: Optional[int] = None
        self.engine = "openpyxl"
        self._source: Any = None
        self._rows: Any = None
        self._next_row = 1

    # -- context management ---------------------------------------------

    def __enter__(self) -> "DataStream":
        self.engine = workbook_engine()
        if not self.sheet:
            self.sheet = pick_sheet(self.path)
            if not self.sheet:
                raise ValueError(f"workbook {self.path} has no worksheets")
        self._source = open_sheet(self.path, self.sheet)
        self.max_row = self._source.max_row
        self.max_column = self._source.max_column

        iterator = self._source.rows()
        buffered: List[Tuple[Any, ...]] = []
        header_values: Optional[Tuple[Any, ...]] = None
        row_number = 0
        for values in iterator:
            buffered.append(values)
            row_number += 1
            if self.header_row is not None:
                if row_number == self.header_row:
                    header_values = values
                    break
            else:
                tokens = {normalise_header(value) for value in values}
                if "symbol" in tokens and "date" in tokens:
                    header_values = values
                    self.header_row = row_number
                    break
                if row_number >= self.scan_rows:
                    break

        if header_values is None:
            # No Symbol/Date header inside the scan window (or an explicit header
            # row past the end of the sheet): behave like the old inspector, which
            # fell back to row 1.
            if self.header_row is None:
                self.header_row = 1
            if 0 < self.header_row <= len(buffered):
                header_values = buffered[self.header_row - 1]
            else:
                header_values = ()

        self.width = len(header_values)
        self.headers = _header_cells(header_values)
        # Rows consumed while hunting for the header are replayed, then the
        # worksheet keeps streaming from where it left off.
        self._rows = chain(buffered[self.header_row:], iterator)
        self._next_row = self.header_row + 1
        return self

    def __exit__(self, *_exc) -> None:
        self.close()

    def close(self) -> None:
        if self._source is not None:
            try:
                self._source.close()
            finally:
                self._source = None

    # -- iteration -------------------------------------------------------

    def rows(self) -> Iterator[Tuple[int, Tuple[Any, ...]]]:
        """Yield ``(row_number, values)`` for every row after the header."""
        for values in self._rows:
            row_number = self._next_row
            self._next_row += 1
            yield row_number, _fit_row(values, self.width)

    def records(self, headers: Optional[List[str]] = None) -> Iterator[Dict[str, Any]]:
        """Yield ``{header: value, "__row__": row_number}`` for non-blank rows."""
        headers = list(headers) if headers is not None else self.headers
        width = len(headers)
        for row_number, values in self.rows():
            if is_blank_row(values):
                continue
            record = {headers[index]: values[index] for index in range(width)}
            record["__row__"] = row_number
            yield record



# --------------------------------------------------------------------------
# MS-OVBA (VBA) decompression — text only, never executed
# --------------------------------------------------------------------------


def _copy_token_help(decompressed_current: int, decompressed_chunk_start: int) -> Tuple[int, int, int]:
    """MS-OVBA 2.4.1.3.19.1 — dynamic CopyToken bit layout.

    The split between the length field and the offset field depends on how far
    into the current chunk the decompressor has already written, which is why a
    fixed ``0x0FFF`` mask cannot decode a real module.
    """
    difference = decompressed_current - decompressed_chunk_start
    bit_count = int(math.ceil(math.log(difference, 2))) if difference > 1 else 0
    bit_count = max(bit_count, 4)
    length_mask = 0xFFFF >> bit_count
    offset_mask = (~length_mask) & 0xFFFF
    return length_mask, offset_mask, bit_count


def decompress_vba_container(data: bytes) -> bytes:
    """Decompress one MS-OVBA ``CompressedContainer`` (MS-OVBA 2.4.1).

    ``data`` must include the leading ``0x01`` signature byte.  The container is
    decoded purely as text — nothing in this module can execute VBA.
    """
    if not data:
        return b""
    container = bytearray(data)
    if container[0] != 0x01:
        raise ValueError(f"invalid MS-OVBA signature byte 0x{container[0]:02X}")

    out = bytearray()
    current = 1
    total = len(container)
    while current + 1 < total:
        chunk_start = current
        header = container[chunk_start] | (container[chunk_start + 1] << 8)
        chunk_size = (header & 0x0FFF) + 3
        compressed = bool((header >> 15) & 0x01)
        end = min(total, chunk_start + chunk_size)
        current = chunk_start + 2

        if not compressed:
            out.extend(container[current:current + 4096])
            current += 4096
            continue

        chunk_out_start = len(out)
        while current < end:
            flag = container[current]
            current += 1
            for bit in range(8):
                if current >= end:
                    break
                if (flag >> bit) & 1:
                    if current + 1 >= end:
                        current = end
                        break
                    token = container[current] | (container[current + 1] << 8)
                    current += 2
                    length_mask, offset_mask, bit_count = _copy_token_help(
                        len(out), chunk_out_start)
                    length = (token & length_mask) + 3
                    offset = ((token & offset_mask) >> (16 - bit_count)) + 1
                    source = len(out) - offset
                    if source < 0:
                        current = end
                        break
                    for index in range(source, source + length):
                        out.append(out[index])
                else:
                    out.append(container[current])
                    current += 1
    return bytes(out)


def extract_vba_modules(xlsm_path: Path) -> Dict[str, str]:
    """Return ``{module_name: source_text}`` for every VBA module found.

    The macros are decompressed as text for documentation purposes only; they
    are never executed by this application.  ``oletools`` is used when it is
    installed because it is a validated implementation of the container format;
    the built-in decompressor above is the fallback.
    """
    modules: Dict[str, str] = {}

    try:
        from oletools.olevba import VBA_Parser  # type: ignore

        parser = VBA_Parser(str(xlsm_path))
        try:
            if parser.detect_vba_macros():
                for _fname, _stream, module_name, code in parser.extract_macros():
                    modules[str(module_name)] = str(code)
        finally:
            parser.close()
        if modules:
            return modules
    except Exception:  # pragma: no cover - oletools missing or malformed file
        modules = {}

    try:
        with zipfile.ZipFile(xlsm_path) as archive:
            if "xl/vbaProject.bin" not in archive.namelist():
                return modules
            blob = archive.read("xl/vbaProject.bin")
    except (zipfile.BadZipFile, OSError):
        return modules

    try:
        import olefile
    except ImportError:  # pragma: no cover - optional dependency
        return modules

    try:
        with olefile.OleFileIO(io.BytesIO(blob)) as ole:
            for entry in ole.listdir():
                if len(entry) != 2 or entry[0] != "VBA":
                    continue
                name = entry[1]
                if name in ("_VBA_PROJECT", "dir"):
                    continue
                raw = ole.openstream(entry).read()
                try:
                    decoded = decompress_vba_container(raw)
                except Exception:  # pragma: no cover - malformed container
                    continue
                if decoded:
                    modules[name] = decoded.decode("latin-1", "replace")
    except Exception:  # pragma: no cover - malformed container
        return modules
    return modules


# --------------------------------------------------------------------------
# Inspection result containers
# --------------------------------------------------------------------------


@dataclass
class CellFormula:
    sheet: str
    coordinate: str
    formula: str
    cached_value: Any = None
    row: int = 0
    column: int = 0

    @property
    def has_cached_value(self) -> bool:
        return self.cached_value is not None

    @property
    def cached_is_error(self) -> bool:
        return isinstance(self.cached_value, str) and self.cached_value.startswith("#")


@dataclass
class SheetInfo:
    name: str
    state: str
    dimensions: str
    max_row: int
    max_column: int
    freeze_panes: Optional[str] = None
    hidden_rows: List[int] = field(default_factory=list)
    hidden_columns: List[str] = field(default_factory=list)
    tables: Dict[str, str] = field(default_factory=dict)
    merged: List[str] = field(default_factory=list)
    data_validations: List[Dict[str, Any]] = field(default_factory=list)
    conditional_formats: List[str] = field(default_factory=list)
    comments: List[Dict[str, str]] = field(default_factory=list)
    hyperlinks: int = 0
    chart_count: int = 0


@dataclass
class WorkbookInspection:
    path: str
    sheets: List[SheetInfo]
    defined_names: List[Dict[str, Any]]
    formulas: List[CellFormula]
    vba_modules: Dict[str, str]
    drawings: List[str]

    def formulas_for(self, sheet: str) -> List[CellFormula]:
        return [f for f in self.formulas if f.sheet == sheet]

    @property
    def sheet_names(self) -> List[str]:
        return [s.name for s in self.sheets]


# --------------------------------------------------------------------------
# Inspector
# --------------------------------------------------------------------------


class WorkbookInspector:
    """Structural inspector for ``.xlsx`` / ``.xlsm`` workbooks."""

    def __init__(self, path: str | Path, read_only: bool = False):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"Workbook not found: {self.path}")
        self._read_only = read_only

    # -- public API ------------------------------------------------------

    def inspect(self, max_formula_rows: Optional[int] = None) -> WorkbookInspection:
        wb_f = openpyxl.load_workbook(self.path, data_only=False, keep_vba=False)
        wb_v = openpyxl.load_workbook(self.path, data_only=True, keep_vba=False)

        sheets: List[SheetInfo] = []
        formulas: List[CellFormula] = []

        for ws in wb_f.worksheets:
            ws_v = wb_v[ws.title]
            info = SheetInfo(
                name=ws.title,
                state=ws.sheet_state,
                dimensions=ws.dimensions,
                max_row=ws.max_row,
                max_column=ws.max_column,
                freeze_panes=str(ws.freeze_panes) if ws.freeze_panes else None,
            )
            info.hidden_rows = [r for r, dim in ws.row_dimensions.items() if dim.hidden]
            info.hidden_columns = [c for c, dim in ws.column_dimensions.items() if dim.hidden]
            try:
                info.tables = {}
                for name in ws.tables:
                    table = ws.tables[name]
                    info.tables[str(name)] = str(getattr(table, "ref", table))
            except Exception:
                info.tables = {}
            info.merged = [str(rng) for rng in ws.merged_cells.ranges]
            try:
                info.data_validations = [
                    {"type": dv.type, "range": str(dv.sqref), "formula1": dv.formula1}
                    for dv in ws.data_validations.dataValidation
                ]
            except Exception:
                info.data_validations = []
            try:
                info.conditional_formats = [str(r.sqref) for r in ws.conditional_formatting]
            except Exception:
                info.conditional_formats = []
            info.comments = [
                {"cell": cell.coordinate, "author": cell.comment.author or "", "text": cell.comment.text}
                for row in ws.iter_rows()
                for cell in row
                if cell.comment is not None
            ]
            info.hyperlinks = len(ws._hyperlinks) if hasattr(ws, "_hyperlinks") else 0
            info.chart_count = len(ws._charts) if hasattr(ws, "_charts") else 0
            sheets.append(info)

            limit = max_formula_rows or ws.max_row
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, limit)):
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        cached = ws_v[cell.coordinate].value
                        formulas.append(
                            CellFormula(
                                sheet=ws.title,
                                coordinate=cell.coordinate,
                                formula=value,
                                cached_value=cached,
                                row=cell.row,
                                column=cell.column,
                            )
                        )

        defined_names = [
            {
                "name": dn.name,
                "scope": dn.localSheetId,
                "hidden": bool(dn.hidden),
                "refers_to": dn.attr_text,
            }
            for dn in wb_f.defined_names.values()
        ]

        return WorkbookInspection(
            path=str(self.path),
            sheets=sheets,
            defined_names=defined_names,
            formulas=formulas,
            vba_modules=extract_vba_modules(self.path),
            drawings=self._list_drawings(),
        )

    def read_data_records(
        self,
        sheet: Optional[str] = None,
        header_row: Optional[int] = None,
    ) -> Tuple[List[str], List[Dict[str, Any]], List[Dict[str, Any]]]:
        """Read the raw market rows from a worksheet.

        Returns ``(headers, records, issues)`` where each record is a dict keyed
        by the *original* header text (nothing is silently dropped).

        The worksheet is streamed (see :class:`DataStream`); for millions of rows
        prefer :meth:`stream_records`, which does not materialise the list.
        """
        headers: List[str] = []
        records: List[Dict[str, Any]] = []
        issues: List[Dict[str, Any]] = []
        resolved_header_row = header_row or 1
        with DataStream(self.path, sheet=sheet, header_row=header_row) as stream:
            headers = list(stream.headers)
            resolved_header_row = stream.header_row or 1
            for record in stream.records(headers):
                records.append(record)

        if not records:
            issues.append({"row": resolved_header_row, "issue": "worksheet contains no data rows"})
        return headers, records, issues

    def sheet_names(self) -> List[str]:
        return sheet_names(self.path)

    # -- internals -------------------------------------------------------

    def _list_drawings(self) -> List[str]:
        try:
            with zipfile.ZipFile(self.path) as archive:
                return [n for n in archive.namelist() if n.startswith("xl/drawings/")]
        except (zipfile.BadZipFile, OSError):
            return []

    def _guess_data_sheet(self, wb) -> Any:
        # Retained for callers that already hold a fully loaded workbook (the
        # non-streaming inspector path).  DataStream has its own read-only
        # equivalent.
        for candidate in ("Data", "data", "Sheet1"):
            if candidate in wb.sheetnames:
                return wb[candidate]
        for ws in wb.worksheets:
            max_column = min(ws.max_column, 12)
            for col in range(1, max_column + 1):
                if normalise_header(ws.cell(1, col).value) in ("symbol", "date"):
                    return ws
        return wb.worksheets[0] if wb.worksheets else None

    def _find_header_row(self, ws) -> int:
        max_row, max_column = min(ws.max_row, 25), ws.max_column
        for row_idx in range(1, max_row + 1):
            tokens = [normalise_header(ws.cell(row_idx, c).value) for c in range(1, max_column + 1)]
            if "symbol" in tokens and "date" in tokens:
                return row_idx
        return 1


__all__ = [
    "DATA_COLUMNS",
    "DATE_COLUMN",
    "SYMBOL_COLUMN",
    "DataStream",
    "PREFERRED_SHEETS",
    "is_blank_row",
    "open_sheet",
    "pick_sheet",
    "sheet_names",
    "workbook_engine",
    "WorkbookInspector",
    "WorkbookInspection",
    "SheetInfo",
    "CellFormula",
    "coerce_date",
    "coerce_number",
    "coerce_symbol",
    "normalise_header",
    "match_canonical_column",
    "extract_vba_modules",
    "decompress_vba_container",
]

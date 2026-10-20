"""Formula parity harness.

The supplied workbook ships with Excel's *cached* results for every formula,
which makes it a ready-made oracle: we recompute the same metrics with
:mod:`services.analytics` and compare cell by cell.

Two independent oracle scopes are exercised:

``analysis``
    ``Analysis!M15:W200`` — the per-session calculated columns — plus the
    ``I2:N6`` lookback summary block and the ``row 13`` baselines.
``main``
    ``Main!B18:J842`` — the screener aggregates and classifications, replayed
    with ``Main!B1`` as the selected date and the sheet's own symbol order (which
    is required to reproduce the ``D{r+38}`` defect faithfully).

Comparison rules
----------------

* Numbers match within an absolute-plus-relative tolerance (default ``1e-6``).
* Cached **error** results (``#VALUE!``, ``#DIV/0!``, ``#REF!`` …) are the
  workbook's way of saying "unavailable"; the application represents those as
  blanks/``None``.  Both sides are normalised accordingly before comparing, and
  the mismatch report records how many cells were unavailable.
* A blank cached string and a cached empty string are the same thing.

Run it from the command line with ``python scripts/parity_check.py``.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import pandas as pd

from .analytics import (
    AnalyticsParams,
    MODE_ORIGINAL,
    compute_lookback_summary,
    compute_screener,
    compute_symbol_metrics,
)

DEFAULT_TOLERANCE = 1e-6

ANALYSIS_FIRST_ROW = 15
ANALYSIS_LAST_ROW = 200
ANALYSIS_HEADER_ROW = 14

MAIN_FIRST_ROW = 18
MAIN_LAST_ROW = 842

#: canonical column -> Analysis worksheet column index
ANALYSIS_SOURCE_COLUMNS = {
    "symbol": 1, "trade_date": 2, "open": 3, "high": 4, "low": 5, "close": 6,
    "last": 7, "prevclose": 8, "tottrdqty": 9, "tottrdval": 10,
    "totaltrades": 11, "delivery_pct": 12,
}

#: Analysis calculated column -> (application column, kind)
ANALYSIS_CALCULATED = {
    "M": ("deliverable_qty", "number"),
    "N": ("avg_price", "number"),
    "O": ("pct_change", "number"),
    "P": ("action", "number"),
    "Q": ("vol_ma", "number"),
    "R": ("vol_level", "string"),
    "S": ("v1", "string"),
    "T": ("badge_bronze", "string"),
    "U": ("badge_silver", "string"),
    "V": ("badge_gold", "string"),
    "W": ("trade_indicator", "string"),
}

ANALYSIS_COLUMN_INDEX = {
    "M": 13, "N": 14, "O": 15, "P": 16, "Q": 17, "R": 18,
    "S": 19, "T": 20, "U": 21, "V": 22, "W": 23,
    "I": 9, "J": 10, "K": 11, "L": 12,
}

#: Analysis summary block row -> lookback window label
SUMMARY_ROWS = {2: 1, 3: 3, 4: 5, 5: 7, 6: 10}
SUMMARY_COLUMNS = {"I": "cmp", "J": "avg_price", "K": "pct_change",
                   "L": "delivery_pct", "M": "action", "N": "cmp_verdict"}
SUMMARY_COLUMN_INDEX = {"I": 9, "J": 10, "K": 11, "L": 12, "M": 13, "N": 14, "P": 16}

MAIN_SOURCE_COLUMN = 1
MAIN_CALCULATED = {
    "B": ("today_delivery", "number"),
    "C": ("avg_delivery", "number"),
    "D": ("today_volume", "number"),
    "E": ("avg_volume", "number"),
    "F": ("today_trades", "number"),
    "G": ("avg_trades", "number"),
    "H": ("d_ind", "string"),
    "I": ("dv1", "string"),
    "J": ("dv2", "string"),
}
MAIN_COLUMN_INDEX = {"B": 2, "C": 3, "D": 4, "E": 5, "F": 6, "G": 7, "H": 8, "I": 9, "J": 10}


# --------------------------------------------------------------------------
# Result containers
# --------------------------------------------------------------------------


@dataclass
class Mismatch:
    cell: str
    expected: Any
    actual: Any
    delta: Optional[float] = None
    kind: str = "value"

    def as_dict(self) -> Dict[str, Any]:
        return {"cell": self.cell, "expected": self.expected, "actual": self.actual,
                "delta": self.delta, "kind": self.kind}


@dataclass
class ParityReport:
    name: str
    mode: str
    tolerance: float
    compared: int = 0
    matched: int = 0
    unavailable: int = 0
    mismatches: List[Mismatch] = field(default_factory=list)
    notes: List[str] = field(default_factory=list)

    @property
    def failed(self) -> int:
        return len(self.mismatches)

    @property
    def passed(self) -> bool:
        return not self.mismatches

    @property
    def match_rate(self) -> float:
        return (self.matched / self.compared) if self.compared else 1.0

    def summary(self) -> Dict[str, Any]:
        return {
            "scope": self.name,
            "mode": self.mode,
            "compared": self.compared,
            "matched": self.matched,
            "unavailable": self.unavailable,
            "failed": self.failed,
            "match_rate": round(self.match_rate, 6),
            "passed": self.passed,
        }

    def mismatch_frame(self) -> pd.DataFrame:
        if not self.mismatches:
            return pd.DataFrame(columns=["cell", "expected", "actual", "delta", "kind"])
        return pd.DataFrame([m.as_dict() for m in self.mismatches])


# --------------------------------------------------------------------------
# Comparison primitives
# --------------------------------------------------------------------------


def _is_error(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("#")


def _normalise_excel(value: Any, kind: str) -> Tuple[Any, bool]:
    """Return ``(comparable_value, unavailable)``."""
    if _is_error(value):
        return ("" if kind == "string" else None), True
    if value is None:
        return ("" if kind == "string" else None), False
    if kind == "string":
        return str(value).strip(), False
    try:
        return float(value), False
    except (TypeError, ValueError):
        return None, True


def _normalise_ours(value: Any, kind: str) -> Any:
    if value is None:
        return "" if kind == "string" else None
    if kind == "string":
        return str(value).strip()
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _numbers_close(a: float, b: float, tolerance: float) -> Tuple[bool, float]:
    delta = abs(a - b)
    scale = max(abs(a), abs(b), 1.0)
    return delta <= tolerance * scale, delta


def _compare(report: ParityReport, cell: str, expected: Any, actual: Any, kind: str) -> None:
    exp, unavailable = _normalise_excel(expected, kind)
    act = _normalise_ours(actual, kind)
    report.compared += 1
    if unavailable:
        report.unavailable += 1

    if kind == "string":
        if exp == act:
            report.matched += 1
        else:
            report.mismatches.append(Mismatch(cell, exp, act, None, "string"))
        return

    if exp is None and act is None:
        report.matched += 1
        return
    if exp is None or act is None:
        report.mismatches.append(Mismatch(cell, exp, act, None, "availability"))
        return
    ok, delta = _numbers_close(exp, act, report.tolerance)
    if ok:
        report.matched += 1
    else:
        report.mismatches.append(Mismatch(cell, exp, act, delta, "number"))


# --------------------------------------------------------------------------
# Oracle readers
# --------------------------------------------------------------------------


def analysis_source_frame(workbook_path: str | Path) -> Tuple[pd.DataFrame, Dict[str, Any], Dict[int, int]]:
    """Read the Analysis block, its cached calculated columns and its formulas.

    Cells whose *formula* contains ``#REF!`` are dropped from the oracle: the
    workbook has destroyed its own reference there (finding B10), so there is
    nothing meaningful to compare against.
    """
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    wbf = openpyxl.load_workbook(workbook_path, data_only=False)
    ws = wb["Analysis"]
    wsf = wbf["Analysis"]
    records: List[Dict[str, Any]] = []
    cached: Dict[str, Any] = {}
    row_map: Dict[int, int] = {}
    for row in range(ANALYSIS_FIRST_ROW, ANALYSIS_LAST_ROW + 1):
        symbol = ws.cell(row, ANALYSIS_SOURCE_COLUMNS["symbol"]).value
        if symbol is None:
            continue
        record = {}
        for name, column in ANALYSIS_SOURCE_COLUMNS.items():
            value = ws.cell(row, column).value
            if name == "trade_date" and isinstance(value, _dt.datetime):
                value = value.date()
            record[name] = value
        records.append(record)
        row_map[row] = len(records) - 1
        for letter, index in ANALYSIS_COLUMN_INDEX.items():
            formula = wsf.cell(row, index).value
            if isinstance(formula, str) and "#REF!" in formula:
                continue
            cached[f"{letter}{row}"] = ws.cell(row, index).value
    wb.close()
    wbf.close()
    return pd.DataFrame.from_records(records), cached, row_map


def summary_oracle(workbook_path: str | Path) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb["Analysis"]
    out: Dict[str, Any] = {}
    for row in SUMMARY_ROWS:
        for letter in SUMMARY_COLUMNS:
            out[f"{letter}{row}"] = ws.cell(row, SUMMARY_COLUMN_INDEX[letter]).value
    for letter in ("L", "P"):
        out[f"{letter}13"] = ws.cell(13, SUMMARY_COLUMN_INDEX[letter]).value
    wb.close()
    return out


def main_oracle(workbook_path: str | Path) -> Tuple[List[str], Dict[str, Any], Any]:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb["Main"]
    symbols: List[str] = []
    cached: Dict[str, Any] = {}
    for row in range(MAIN_FIRST_ROW, MAIN_LAST_ROW + 1):
        symbol = ws.cell(row, MAIN_SOURCE_COLUMN).value
        if symbol is None:
            continue
        symbols.append(str(symbol))
        for letter, index in MAIN_COLUMN_INDEX.items():
            cached[f"{letter}{row}"] = ws.cell(row, index).value
    as_of = ws.cell(1, 2).value
    if isinstance(as_of, _dt.datetime):
        as_of = as_of.date()
    wb.close()
    return symbols, cached, as_of


# --------------------------------------------------------------------------
# Comparisons
# --------------------------------------------------------------------------


def compare_analysis(workbook_path: str | Path, mode: str = MODE_ORIGINAL,
                     tolerance: float = DEFAULT_TOLERANCE,
                     limit_mismatches: int = 40) -> ParityReport:
    report = ParityReport("analysis", mode, tolerance)
    frame, cached, row_map = analysis_source_frame(workbook_path)
    if frame.empty:
        report.notes.append("the Analysis block is empty — nothing to compare")
        return report

    params = AnalyticsParams(mode=mode)
    metrics = compute_symbol_metrics(frame, params)

    for cell, expected in cached.items():
        letter = cell[0]
        row = int(cell[1:])
        if letter not in ANALYSIS_CALCULATED or row not in row_map:
            continue
        column, kind = ANALYSIS_CALCULATED[letter]
        index = row_map[row]
        actual = metrics[column].iloc[index]
        if pd.isna(actual):
            actual = None
        _compare(report, cell, expected, actual, kind)
        if len(report.mismatches) >= limit_mismatches:
            report.notes.append(f"stopped after {limit_mismatches} mismatches")
            break

    if len(report.mismatches) < limit_mismatches:
        oracle = summary_oracle(workbook_path)
        summary = compute_lookback_summary(frame, params, reference="earliest")
        by_window = {entry["window"]: entry for entry in summary}
        for row, window in SUMMARY_ROWS.items():
            entry = by_window.get(f"{window}D")
            if entry is None:
                continue
            for letter, name in SUMMARY_COLUMNS.items():
                kind = "string" if letter == "N" else "number"
                _compare(report, f"{letter}{row}", oracle.get(f"{letter}{row}"),
                         entry.get(name), kind)

        for cell, attribute in (("L13", "baseline_delivery"), ("P13", "baseline_action")):
            _compare(report, cell, oracle.get(cell), metrics.attrs.get(attribute), "number")

    report.notes.append(
        f"{report.unavailable} compared cells are unavailable in the workbook "
        f"(cached error results), represented by blank values here"
    )
    return report


def compare_screener(workbook_path: str | Path, data_frame: pd.DataFrame,
                     mode: str = MODE_ORIGINAL,
                     tolerance: float = DEFAULT_TOLERANCE,
                     limit_mismatches: int = 40) -> ParityReport:
    report = ParityReport("main", mode, tolerance)
    symbols, cached, as_of = main_oracle(workbook_path)
    params = AnalyticsParams(mode=mode, symbol_order=symbols)
    result = compute_screener(data_frame, as_of=as_of, params=params, universe=symbols)

    if result.empty:
        report.notes.append("the screener produced no rows — nothing to compare")
        return report

    indexed = result.set_index("symbol")
    for row in range(MAIN_FIRST_ROW, MAIN_LAST_ROW + 1):
        symbol = symbols[row - MAIN_FIRST_ROW]
        if symbol not in indexed.index:
            continue
        record = indexed.loc[symbol]
        for letter, (column, kind) in MAIN_CALCULATED.items():
            cell = f"{letter}{row}"
            expected = cached.get(cell)
            actual = record[column]
            if pd.isna(actual):
                actual = None
            _compare(report, cell, expected, actual, kind)
            if len(report.mismatches) >= limit_mismatches:
                report.notes.append(f"stopped after {limit_mismatches} mismatches")
                return report

    report.notes.append(
        f"replayed with Main!B1 = {as_of} and the sheet's own symbol order so that "
        f"the D{{r+38}} GOLD test of finding B5 is reproduced exactly"
    )
    return report


def run_all(workbook_path: str | Path, data_frame: Optional[pd.DataFrame] = None,
            mode: str = MODE_ORIGINAL,
            tolerance: float = DEFAULT_TOLERANCE) -> List[ParityReport]:
    reports = [compare_analysis(workbook_path, mode=mode, tolerance=tolerance)]
    if data_frame is not None and not data_frame.empty:
        reports.append(compare_screener(workbook_path, data_frame, mode=mode,
                                        tolerance=tolerance))
    return reports


__all__ = [
    "ParityReport",
    "Mismatch",
    "compare_analysis",
    "compare_screener",
    "run_all",
    "analysis_source_frame",
    "main_oracle",
    "summary_oracle",
    "DEFAULT_TOLERANCE",
]

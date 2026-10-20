#!/usr/bin/env python3
"""Regenerate ``config/palette_workbook.json`` from the sample workbook.

The dashboard is painted by conditional formatting, so the colours live in two
places: ``xl/styles.xml`` (the dxf table) and the ``<conditionalFormatting>``
rules on ``Main`` / ``Analysis``.  This script reads both and writes the JSON
the application ships.

Usage::

    python scripts/extract_palette.py
    python scripts/extract_palette.py --workbook path/to/file.xlsm --out /tmp/p.json

The chart colourway is a hand-picked subset of the same palette (the workbook
itself has no charts), so it is kept as a constant here.
"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, Optional
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = ROOT / "data" / "Saif-Stock Market Data File-Sample.xlsm"
DEFAULT_OUT = ROOT / "config" / "palette_workbook.json"

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

#: Which sheet column feeds which palette group.  ``H``/``I``/``J`` are the
#: screener's D.Ind / DV1 / DV2 columns on ``Main``; the range colours come
#: from the "Above / Below" rules on ``Analysis``.
GROUP_COLUMNS = {
    "d_ind": ("sheet3.xml", "H"),
    "dv1": ("sheet3.xml", "I"),
    "dv2": ("sheet3.xml", "J"),
    "range": ("sheet2.xml", "N"),
}

SERIES = [
    "#4472C4", "#FFC000", "#00B050", "#7030A0", "#00B0F0",
    "#BF9000", "#92D050", "#AFABAB", "#FF0000", "#767171",
]

NOTES = [
    "Main!A1 = #FFFF00 date label, Main!B1 = #9DC3E6 date value.",
    "Main!N18:N29 is the legend; the screener fills come from conditional-formatting dxfs 151-174.",
    "Analysis!A14:X14 uses #FFD966 for the price columns and #767171 for the derived columns.",
    "Analysis!N/R/W use #C6EFCE/#006100 (above) and #FFC7CE/#9C0006 (below).",
    "SILVER is #AFABAB on Main but #A6A6A6 on Analysis; the Main dashboard value wins here.",
    "The legend's Low Del swatch is #FEE2E3 while the rule paints #FDD7D5; the rule wins here.",
    "V1 (J/Spl/Trpl/Dbl/--/^) is uncoloured in the workbook, so the entries are null.",
]


def _rgb(element: Optional[ET.Element]) -> Optional[str]:
    if element is None:
        return None
    value = element.get("rgb")
    if value and len(value) == 8 and value[:2].upper() == "FF":
        value = value[2:]
    return "#" + value.upper() if value else None


def _dxf_table(archive: zipfile.ZipFile) -> list:
    root = ET.fromstring(archive.read("xl/styles.xml"))
    dxfs = root.find("m:dxfs", NS)
    return list(dxfs) if dxfs is not None else []


def _dxf_colours(dxf: ET.Element) -> tuple:
    pattern = dxf.find("m:fill/m:patternFill", NS)
    fill = None
    if pattern is not None:
        fill = _rgb(pattern.find("m:fgColor", NS)) or _rgb(pattern.find("m:bgColor", NS))
    font = dxf.find("m:font", NS)
    text = _rgb(font.find("m:color", NS)) if font is not None else None
    return fill, text


def _conditional_rules(archive: zipfile.ZipFile, sheet: str) -> list:
    xml = archive.read(f"xl/worksheets/{sheet}").decode("utf-8", "replace")
    rules = []
    for block in re.finditer(
            r'<conditionalFormatting sqref="([^"]+)">(.*?)</conditionalFormatting>',
            xml, re.S):
        sqref, body = block.group(1), block.group(2)
        for rule in re.finditer(r'<cfRule[^>]*dxfId="(\d+)"[^>]*>(.*?)</cfRule>',
                                body, re.S):
            label = re.search(r"&quot;([^&]+)&quot;", rule.group(2))
            if not label:
                continue
            rules.append((sqref, int(rule.group(1)), label.group(1)))
    return rules


def _first_column(sqref: str) -> str:
    return re.match(r"([A-Z]+)", sqref).group(1)


def extract(workbook: Path) -> Dict[str, Any]:
    archive = zipfile.ZipFile(workbook)
    dxfs = _dxf_table(archive)

    groups: Dict[str, Dict[str, Any]] = {}
    for group, (sheet, column) in GROUP_COLUMNS.items():
        entries: Dict[str, Any] = {}
        for sqref, dxf_id, label in _conditional_rules(archive, sheet):
            if _first_column(sqref) != column:
                continue
            fill, text = _dxf_colours(dxfs[dxf_id])
            if not fill:
                continue
            entry: Dict[str, Any] = {"fill": fill}
            if text and text.upper() != "#FFFFFF":
                entry["text"] = text
            entries.setdefault(label, entry)
        groups[group] = entries

    # V1 is genuinely uncoloured; keep the labels so the editor can offer them.
    groups["v1"] = {label: None for label in ("J", "Spl", "Trpl", "Dbl", "--", "^")}

    sheet = load_workbook(workbook, data_only=True)["Main"]
    analysis = load_workbook(workbook, data_only=True)["Analysis"]

    def cell_fill(worksheet, coordinate: str) -> Optional[str]:
        fill = worksheet[coordinate].fill
        if fill is None or not fill.fill_type:
            return None
        colour = fill.fgColor
        return "#" + colour.rgb[-6:].upper() if colour.type == "rgb" else None

    chrome = {
        "surface": "#FFFFFF",
        "table_header": cell_fill(analysis, "A14"),
        "table_header_text": "#101828",
        "section": cell_fill(analysis, "M14"),
        "accent": cell_fill(analysis, "G1"),
        "highlight": cell_fill(analysis, "X13"),
        "date_label": cell_fill(sheet, "A1"),
        "date_value": cell_fill(sheet, "B1"),
    }

    return {
        "version": 1,
        "name": "Saif workbook",
        "source": workbook.name,
        "notes": NOTES,
        "groups": groups,
        "chrome": {key: value for key, value in chrome.items() if value},
        "series": SERIES,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--stdout", action="store_true", help="print instead of writing")
    args = parser.parse_args()

    if not args.workbook.exists():
        raise SystemExit(f"workbook not found: {args.workbook}")
    document = extract(args.workbook)
    payload = json.dumps(document, indent=2) + "\n"
    if args.stdout:
        print(payload, end="")
    else:
        args.out.write_text(payload, encoding="utf-8")
        print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
